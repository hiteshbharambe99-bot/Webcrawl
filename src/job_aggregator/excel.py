from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

from .models import JobListing, SHEET_HEADERS
from .utils import to_sheet_row


class LocalExcelWriter:
    def __init__(self, file_path: str, worksheet_name: str = "Jobs_AU"):
        self.file_path = Path(file_path)
        self.worksheet_name = worksheet_name

    def _load_workbook(self):
        from openpyxl import Workbook, load_workbook

        self.file_path.parent.mkdir(parents=True, exist_ok=True)
        if self.file_path.exists():
            return load_workbook(self.file_path)
        wb = Workbook()
        ws = wb.active
        ws.title = self.worksheet_name
        ws.append(SHEET_HEADERS)
        wb.save(self.file_path)
        return wb

    def _get_sheet(self, wb):
        if self.worksheet_name in wb.sheetnames:
            ws = wb[self.worksheet_name]
        else:
            ws = wb.create_sheet(self.worksheet_name)
            ws.append(SHEET_HEADERS)
        if ws.max_row == 0:
            ws.append(SHEET_HEADERS)
        return ws

    def _load_existing_index(self, ws) -> dict[str, int]:
        index: dict[str, int] = {}
        for row_num in range(2, ws.max_row + 1):
            val = ws.cell(row=row_num, column=15).value
            if val:
                index[str(val)] = row_num
        return index

    def upsert_jobs(self, jobs: Iterable[JobListing]) -> tuple[int, int]:
        wb = self._load_workbook()
        ws = self._get_sheet(wb)
        existing = self._load_existing_index(ws)

        inserted = 0
        updated = 0

        for job in jobs:
            row = to_sheet_row(job)
            target = existing.get(job.canonical_key)
            if target:
                for col_idx, value in enumerate(row, start=1):
                    ws.cell(row=target, column=col_idx, value=value)
                updated += 1
            else:
                ws.append(row)
                inserted += 1

        wb.save(self.file_path)
        return inserted, updated
